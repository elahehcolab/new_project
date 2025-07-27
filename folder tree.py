import os

def get_folder_structure(folder_path, indent=0):
    structure = ""
    for item in sorted(os.listdir(folder_path)):
        path = os.path.join(folder_path, item)
        if os.path.isdir(path):
            structure += '  ' * indent + '|-- ' + item + '\n'
            structure += get_folder_structure(path, indent + 1)
    return structure

# === Set your target folder path here ===
folder_path = r"C:\Users\user\Desktop\Saba"

# === Output file names ===
output_txt = "folder_tree_only_dirs.txt"
output_docx = "folder_tree_only_dirs.docx"
output_pdf = "folder_tree_only_dirs.pdf"

# === Generate structure string ===
structure = get_folder_structure(folder_path)

# === Save as TXT ===
with open(output_txt, "w", encoding="utf-8") as f:
    f.write(structure)

# === Save as DOCX ===
from docx import Document
doc = Document()
doc.add_paragraph(structure)
doc.save(output_docx)

# === Save as PDF ===
from fpdf import FPDF

pdf = FPDF()
pdf.add_page()
pdf.set_font("Courier", size=10)
for line in structure.split('\n'):
    pdf.cell(0, 5, txt=line, ln=1)
pdf.output(output_pdf)

print("Folder structure saved as TXT, DOCX, and PDF.")
